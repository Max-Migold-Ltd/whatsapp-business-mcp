"""
Create the Excel workbook the support bot logs to.

Builds "Complaints Log.xlsx" with:
  - Tickets sheet:        table "Tickets", one column per question in flow.json,
                          Status dropdown (Pending / In Progress / Resolved) with colours
  - Conversations sheet:  table "Conversations", every WhatsApp message
  - How to use sheet:     short instructions for the agent

Columns are read from support_bot/flow.json, so re-run this after changing the
questions. Upload the result to the SharePoint folder set in SHAREPOINT_FILE_PATH.

Usage:
    pip install openpyxl
    python scripts/create_excel_log.py                      # -> Complaints Log.xlsx
    python scripts/create_excel_log.py "My Log.xlsx"        # custom file name
"""

import argparse
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_FLOW = os.path.join(ROOT, "support_bot", "flow.json")

# Must match support_bot/sharepoint.py (EXCEL_COLUMNS and ticket_columns)
CONVERSATION_COLUMNS = [
    "Date/Time", "Phone", "Customer Name", "Direction", "Sent By",
    "Type", "Message", "Conversation State", "WhatsApp Message ID",
]

# Rows covered by the Status dropdown and formatting (tables grow into them)
MAX_ROWS = 20000

STATUS_COLOURS = {
    "Pending": "FCE4B6",      # amber
    "In Progress": "BDD7EE",  # blue
    "Resolved": "C6EFCE",     # green
}

COLUMN_WIDTHS = {
    "Ticket #": 12, "Date/Time": 19, "Phone": 17, "Customer Name": 22, "Direction": 10,
    "Sent By": 10, "Type": 12, "Message": 60, "Complaint": 60, "Conversation State": 12,
    "WhatsApp Message ID": 20, "Status": 14,
}
WRAPPED = {"Message", "Complaint"}


def load_flow(path):
    with open(path, encoding="utf-8") as f:
        flow = json.load(f)
    fields = []
    for step in flow["steps"].values():
        if step["field"] not in fields:
            fields.append(step["field"])
    statuses = flow.get("ticket_statuses") or [flow.get("ticket_status", "Pending")]
    return fields, statuses


def add_table(ws, name, columns, style):
    """Header row + one empty row (Excel tables need at least one data row)."""
    for col, header in enumerate(columns, start=1):
        ws.cell(row=1, column=col, value=header)
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(header, 18)
        if header in WRAPPED:
            for row in range(2, 3):
                ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    ref = f"A1:{get_column_letter(len(columns))}2"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"


def column_range(columns, header):
    letter = get_column_letter(columns.index(header) + 1)
    return f"{letter}2:{letter}{MAX_ROWS}"


def build(output, flow_path):
    fields, statuses = load_flow(flow_path)
    ticket_columns = ["Ticket #", "Date/Time", "Phone", "Customer Name", *fields, "Status"]

    wb = Workbook()

    # ---- Tickets ----
    tickets = wb.active
    tickets.title = "Tickets"
    add_table(tickets, "Tickets", ticket_columns, "TableStyleMedium2")

    status_cells = column_range(ticket_columns, "Status")
    dropdown = DataValidation(
        type="list",
        formula1='"' + ",".join(statuses) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid status",
        error="Choose one of: " + ", ".join(statuses),
    )
    tickets.add_data_validation(dropdown)
    dropdown.add(status_cells)

    for status, colour in STATUS_COLOURS.items():
        if status in statuses:
            tickets.conditional_formatting.add(
                status_cells,
                CellIsRule(operator="equal", formula=[f'"{status}"'],
                           fill=PatternFill("solid", start_color=colour, end_color=colour),
                           font=Font(bold=True)),
            )

    # ---- Conversations ----
    conversations = wb.create_sheet("Conversations")
    add_table(conversations, "Conversations", CONVERSATION_COLUMNS, "TableStyleLight9")

    # ---- How to use ----
    guide = wb.create_sheet("How to use")
    guide.column_dimensions["A"].width = 110
    lines = [
        ("Max Migold Facility Management - WhatsApp Complaints Log", True),
        ("", False),
        ("TICKETS sheet - one row per complaint, added automatically by the WhatsApp bot.", True),
        ("• New complaints arrive with Status = Pending.", False),
        ("• Update the Status using the dropdown: " + " → ".join(statuses) + ".", False),
        ("• Residents see this status when they send their ticket ID (e.g. MMF-00042) on WhatsApp.", False),
        ("• You may sort, filter and add notes in a NEW column to the right, but see the rules below.", False),
        ("", False),
        ("CONVERSATIONS sheet - every WhatsApp message (resident, bot and agent), added automatically.", True),
        ("", False),
        ("RULES - so the bot keeps working", True),
        ("• Do NOT rename, move or delete the column headers.", False),
        ("• Do NOT rename the tables ('Tickets', 'Conversations') or the file, and do not move the file.", False),
        ("• Do NOT add columns inside the tables - the bot writes an exact number of columns.", False),
        ("• Do NOT edit the Ticket # column.", False),
        ("• The first empty row in each table is normal - leave it.", False),
    ]
    for row, (text, bold) in enumerate(lines, start=1):
        cell = guide.cell(row=row, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if row == 1 else 11)

    wb.save(output)
    return ticket_columns


def main():
    parser = argparse.ArgumentParser(description="Create the Excel workbook for the WhatsApp support bot.")
    parser.add_argument("output", nargs="?", default="Complaints Log.xlsx", help="file to create")
    parser.add_argument("--flow", default=DEFAULT_FLOW, help="path to flow.json")
    args = parser.parse_args()

    if os.path.exists(args.output):
        sys.exit(f"'{args.output}' already exists - delete it or choose another name.")

    columns = build(args.output, args.flow)
    print(f"Created {args.output}")
    print("Tickets columns: " + " | ".join(columns))
    print("Next: upload it to SharePoint and set SHAREPOINT_FILE_PATH to its path in the Documents library.")


if __name__ == "__main__":
    main()
